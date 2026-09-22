from django.shortcuts import get_object_or_404, render, redirect
from django.http import HttpResponse,JsonResponse, HttpResponseRedirect
from django.urls import reverse, reverse_lazy
from userauths.models import *
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
# Create your views here.
from django.contrib import messages
from datetime import datetime
from .utils import (
    send_email_with_html_body,
    send_welcome_activation_email,
    send_verification_email,
    send_password_reset_otp_email,
    send_account_activation_email,
)
from userauths.forms import *
from django.views.generic import CreateView, ListView,UpdateView, DetailView,DeleteView,TemplateView
from django.contrib.auth import login, authenticate, logout
from django.contrib.auth.mixins import LoginRequiredMixin
from django.utils import timezone
from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import make_password
from django.conf import settings
from django.views import View
from django.contrib.auth.models import User
import random
from django.core.mail import EmailMultiAlternatives
from django.template.loader import render_to_string
from django.db.models import Q, Sum
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.db import transaction
from decimal import Decimal
from xhtml2pdf import pisa
import os
import tempfile
import string
import random
import json

from .forms import (
    ChangePasswordForm,
    ClientRegistrationForm,
    PasswordResetRequestForm,
    PasswordResetConfirmForm,
    ResendActivationForm,
    ProfileSelfUserForm,
    ProfileSelfProfilForm,
)
from .decorators import superuser_required
from django.contrib.auth.views import PasswordChangeView
from django.contrib.auth import update_session_auth_hash
from .models import PWD_FORGET, EmailVerificationToken, ProfilUser
from .context_processors import _user_nav_initials, _user_has_custom_photo

def pbholdingsiteview(request):
    return render(request, 'page/pbholdingsite_home.html',)

def list_users(request):
    users = CustomUser.objects.filter(is_superuser=False)
    return render(request, 'liste_compte.html', {'users': users})

def generate_random_password(length=8):
    characters = string.ascii_letters + string.digits 
    return ''.join(random.choice(characters) for i in range(length))


def _invalidate_activation_tokens(user):
    EmailVerificationToken.objects.filter(user=user, used=False).update(
        used=True,
        used_at=timezone.now(),
    )


def _send_account_activation(user, request, temporary_password=None):
    _invalidate_activation_tokens(user)
    token_obj = EmailVerificationToken.objects.create(user=user)
    return send_welcome_activation_email(
        user,
        token_obj.token,
        request=request,
        temporary_password=temporary_password,
    )


def _redirect_after_login(user, request=None):
    # Guide commande ecom : utile surtout pour les clients
    if request is not None and getattr(user, 'role', None) == 'client':
        request.session['show_ecom_order_guide'] = True

    role = getattr(user, 'role', None)

    if role == 'admin' or getattr(user, 'is_superuser', False):
        return redirect('tbord')
    if role == 'chefagence':
        return redirect('tbord')
    if role == 'livreur':
        return redirect('livraison_cmd_online')
    if role == 'accueil':
        return redirect('paniers')
    if role == 'caissier':
        return redirect('caissiere')
    if role == 'gestionnaire':
        return redirect('stock')
    if role == 'client':
        return redirect(f"{reverse('ecom_index')}?order_guide=1")

    # Fallback (rôles inattendus)
    return redirect('paniers')

# @login_required(login_url='connexion')
def AddCompte(request):
    q = (request.GET.get('q') or '').strip()
    # Comptes staff uniquement — les clients sont gérés sur compte_client.
    # Les superutilisateurs sont masqués : ce sont les comptes techniques
    # d'administration, ils n'ont pas à être modifiés ni désactivés depuis
    # cette page au risque de se verrouiller hors de l'application.
    comptes = CustomUser.objects.select_related('local_entrepot').exclude(
        role='client',
    ).exclude(
        is_superuser=True,
    ).order_by('username')
    if q:
        comptes = comptes.filter(
            Q(username__icontains=q)
            | Q(email__icontains=q)
            | Q(contact__icontains=q)
            | Q(role__icontains=q)
            | Q(genre__icontains=q)
            | Q(local_entrepot__nom__icontains=q)
        )

    if request.method == 'POST':
        user_form = CustomUserForm(request.POST)
        permission_form = UserPermissionForm(request.POST)
        if user_form.is_valid() and permission_form.is_valid():
            try:
                user = user_form.save(commit=False)
                password = generate_random_password()
                user.set_password(password)
                user.is_active = False
                user.save()
                permissions = permission_form.cleaned_data.get('permissions', [])
                if permissions:
                    user.custom_permissions.set(permissions)
                else:
                    user.custom_permissions.clear()
                ProfilUser.objects.get_or_create(user=user)
                has_send = _send_account_activation(user, request, temporary_password=password)
                if has_send:
                    messages.success(
                        request,
                        f"Compte créé. Un email d'activation a été envoyé à {user.email}.",
                    )
                else:
                    messages.warning(
                        request,
                        f"Compte créé, mais l'email d'activation n'a pas pu être envoyé à {user.email}.",
                    )
                return redirect('add_compte')
            except Exception as e:
                messages.error(request, f"Erreur: {str(e)}")
        else:
            for error in user_form.errors.values():
                messages.error(request, error)
    else:
        user_form = CustomUserForm()
        permission_form = UserPermissionForm()
    context = {
        'user_form': user_form,
        'permission_form': permission_form,
        'comptes': comptes,
        'search_q': q,
    }
    return render(request, 'page/add_compte.html', context)


@login_required(login_url='connexion')
def export_comptes_excel(request):
    """Exporter la liste des comptes staff vers Excel (exclut les clients)."""
    q = (request.GET.get('q') or '').strip()
    comptes = CustomUser.objects.select_related('local_entrepot').prefetch_related(
        'custom_permissions',
    ).exclude(role='client').order_by('username')
    if q:
        comptes = comptes.filter(
            Q(username__icontains=q)
            | Q(email__icontains=q)
            | Q(contact__icontains=q)
            | Q(role__icontains=q)
            | Q(genre__icontains=q)
            | Q(local_entrepot__nom__icontains=q)
        )

    if not comptes.exists():
        messages.warning(request, 'Aucun compte à exporter.')
        return redirect('add_compte')

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Comptes'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='EA580C', end_color='EA580C', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    headers = [
        'Nom utilisateur', 'Email', 'Contact', 'Rôle', 'Genre',
        'Entrepôt', 'Actif', 'Permissions',
    ]
    for col_num, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    widths = {'A': 22, 'B': 32, 'C': 16, 'D': 16, 'E': 12, 'F': 22, 'G': 10, 'H': 50}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    role_labels = dict(ROLE_CHOICES)
    for row_num, compte in enumerate(comptes, start=2):
        perms = ', '.join(compte.custom_permissions.values_list('name', flat=True))
        sheet.cell(row=row_num, column=1).value = compte.username
        sheet.cell(row=row_num, column=2).value = compte.email
        sheet.cell(row=row_num, column=3).value = compte.contact or ''
        sheet.cell(row=row_num, column=4).value = role_labels.get(compte.role, compte.role)
        sheet.cell(row=row_num, column=5).value = compte.genre or ''
        sheet.cell(row=row_num, column=6).value = (
            compte.local_entrepot.nom if compte.local_entrepot_id else ''
        )
        sheet.cell(row=row_num, column=7).value = 'Oui' if compte.is_active else 'Non'
        sheet.cell(row=row_num, column=8).value = perms

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'comptes_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


@login_required(login_url='connexion')
def list_comptes_clients(request):
    """Liste des comptes clients (role=client)."""
    q = (request.GET.get('q') or '').strip()
    comptes = CustomUser.objects.select_related('profil').filter(
        role='client',
    ).order_by('username')
    if q:
        comptes = comptes.filter(
            Q(username__icontains=q)
            | Q(email__icontains=q)
            | Q(contact__icontains=q)
            | Q(first_name__icontains=q)
            | Q(last_name__icontains=q)
            | Q(genre__icontains=q)
            | Q(profil__ville__icontains=q)
            | Q(profil__adresse__icontains=q)
        )
    return render(request, 'page/compte_client.html', {
        'comptes': comptes,
        'search_q': q,
    })


def _context_fiche_compte_client(compte):
    """Contexte partagé fiche client (modal + PDF) : profil, stats commandes, favoris."""
    from stock.models import Commande
    from ecommerce.models import FavoriPiece

    profil = getattr(compte, 'profil', None)
    commandes_qs = Commande.objects.filter(
        commande_en_ligne=True,
        panier__utilisateur=compte,
    ).select_related(
        'panier', 'panier__local_entrepot', 'livreur',
    ).order_by('-date')

    payees_qs = commandes_qs.filter(paye=True)
    stats = {
        'nb_commandes': commandes_qs.count(),
        'nb_payees': payees_qs.count(),
        'montant_paye': payees_qs.aggregate(s=Sum('total'))['s'] or Decimal('0'),
        'nb_annulees': commandes_qs.filter(statut_commande='annuler').count(),
        'nb_livrees': commandes_qs.filter(statut_commande='livrer').count(),
        'nb_en_attente': commandes_qs.filter(statut_commande='en_attente').count(),
        'nb_validees': commandes_qs.filter(statut_commande='valider').count(),
    }

    favoris = (
        FavoriPiece.objects.filter(utilisateur=compte)
        .select_related('piece', 'piece__categorie', 'piece__sous_categorie')
        .order_by('-date_ajout')
    )

    return {
        'object': compte,
        'profil': profil,
        'stats': stats,
        'commandes_recentes': list(commandes_qs[:10]),
        'favoris': favoris,
        'nb_favoris': favoris.count(),
    }


@login_required(login_url='connexion')
def detail_compte_client(request, pk):
    """Fiche compte client (fragment modal AJAX)."""
    compte = get_object_or_404(
        CustomUser.objects.select_related('profil'),
        pk=pk,
        role='client',
    )
    return render(
        request,
        'partials/compte_client_detail.html',
        _context_fiche_compte_client(compte),
    )


@login_required(login_url='connexion')
def imprimer_fiche_client_pdf(request, pk):
    """Exporter la fiche client au format PDF."""
    compte = get_object_or_404(
        CustomUser.objects.select_related('profil'),
        pk=pk,
        role='client',
    )
    context = _context_fiche_compte_client(compte)
    context['pdf_export'] = True
    context['generated_at'] = timezone.localtime()
    html_content = render_to_string(
        'page/fiche_client_pdf.html', context, request=request,
    )
    safe_name = ''.join(
        c if c.isalnum() or c in '-_' else '_'
        for c in (compte.username or f'client_{pk}')
    )
    filename = f'fiche_client_{safe_name}.pdf'
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    pisa_status = pisa.CreatePDF(html_content, dest=response)
    if pisa_status.err:
        return HttpResponse('Erreur lors de la génération du PDF.', status=500)
    return response


class UpdateCompteClientView(LoginRequiredMixin, UpdateView):
    login_url = 'connexion'
    model = CustomUser
    form_class = ClientCompteForm
    template_name = 'partials/compte_client_form.html'
    success_url = reverse_lazy('compte_client')
    success_message = 'Compte client modifié avec succès✓✓'

    def get_queryset(self):
        return CustomUser.objects.filter(role='client')

    def form_valid(self, form):
        response = super().form_valid(form)
        if self.request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({
                'success': True,
                'message': self.success_message,
                'redirect_url': reverse('compte_client'),
            })
        messages.success(self.request, self.success_message)
        return response

    def form_invalid(self, form):
        if self.request.headers.get('x-requested-with') == 'XMLHttpRequest':
            html = render_to_string(
                self.template_name,
                {'form': form, 'object': self.object},
                request=self.request,
            )
            return JsonResponse({'success': False, 'html': html})
        return super().form_invalid(form)


@login_required(login_url='connexion')
def delete_compte_client(request, pk):
    """Supprimer un compte client."""
    try:
        user = get_object_or_404(CustomUser, id=pk, role='client')
        username = user.username
        user.delete()
        messages.success(request, f"Compte client de {username} supprimé avec succès.")
    except Exception as e:
        messages.error(request, f"Erreur lors de la suppression : {str(e)}")
    return redirect('compte_client')


@login_required(login_url='connexion')
def export_comptes_clients_excel(request):
    """Exporter la liste des comptes clients vers Excel."""
    q = (request.GET.get('q') or '').strip()
    comptes = CustomUser.objects.select_related('profil').filter(
        role='client',
    ).order_by('username')
    if q:
        comptes = comptes.filter(
            Q(username__icontains=q)
            | Q(email__icontains=q)
            | Q(contact__icontains=q)
            | Q(first_name__icontains=q)
            | Q(last_name__icontains=q)
            | Q(genre__icontains=q)
            | Q(profil__ville__icontains=q)
            | Q(profil__adresse__icontains=q)
        )

    if not comptes.exists():
        messages.warning(request, 'Aucun compte client à exporter.')
        return redirect('compte_client')

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Comptes clients'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='EA580C', end_color='EA580C', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    headers = [
        'Nom utilisateur', 'Prénom', 'Nom', 'Email', 'Contact', 'Genre',
        'Actif', 'Adresse', 'Ville', 'Pays', 'Statut profil',
    ]
    for col_num, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    widths = {
        'A': 22, 'B': 16, 'C': 16, 'D': 32, 'E': 16, 'F': 12,
        'G': 10, 'H': 28, 'I': 16, 'J': 18, 'K': 14,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    for row_num, compte in enumerate(comptes, start=2):
        profil = getattr(compte, 'profil', None)
        sheet.cell(row=row_num, column=1).value = compte.username
        sheet.cell(row=row_num, column=2).value = compte.first_name or ''
        sheet.cell(row=row_num, column=3).value = compte.last_name or ''
        sheet.cell(row=row_num, column=4).value = compte.email
        sheet.cell(row=row_num, column=5).value = compte.contact or ''
        sheet.cell(row=row_num, column=6).value = compte.genre or ''
        sheet.cell(row=row_num, column=7).value = 'Oui' if compte.is_active else 'Non'
        sheet.cell(row=row_num, column=8).value = (profil.adresse if profil else '') or ''
        sheet.cell(row=row_num, column=9).value = (profil.ville if profil else '') or ''
        sheet.cell(row=row_num, column=10).value = (profil.pays if profil else '') or ''
        sheet.cell(row=row_num, column=11).value = (
            profil.get_statut_display() if profil else ''
        )

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'comptes_clients_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


# class UpdateCompteView(LoginRequiredMixin, CustomPermissionRequiredMixin, UpdateView):
class UpdateCompteView(LoginRequiredMixin, UpdateView):
    login_url = 'connexion'
    # permission_url = 'updat_recet'
    model = CustomUser
    form_class = CustomUserForm
    template_name = "partials/compte_form.html"
    success_url = reverse_lazy('add_compte')
    success_message = 'Compte modifiée avec succès✓✓'

    def get_form_kwargs(self):
        """Suffixe les identifiants des champs de la modale.

        La page add_compte affiche deja le formulaire de creation, qui rend
        `id_role` et `id_local_entrepot`. Sans suffixe, la modale chargee en
        AJAX dupliquerait ces identifiants : le script de bascule de la page
        piloterait alors le mauvais champ, et le HTML serait invalide.
        """
        kwargs = super().get_form_kwargs()
        kwargs['auto_id'] = 'id_%s_edit'
        return kwargs

    def form_valid(self, form):
        response = super().form_valid(form)
        if self.request.headers.get("x-requested-with") == "XMLHttpRequest":
            return JsonResponse({
                "success": True,
                "message": self.success_message,
                "redirect_url": reverse("add_compte"),
            })
        return response
    def form_invalid(self, form):
        if self.request.headers.get("x-requested-with") == "XMLHttpRequest":
            html = render_to_string("partials/compte_form.html" , {"form": form, "object": self.object}, request=self.request)
            return JsonResponse({"success": False, "html": html})
        return super().form_invalid(form)

def _auth_context(**extra):
    """Contexte commun aux pages de connexion et d'inscription.

    `google_client_id` conditionne l'affichage du bouton Google : vide, le bloc
    entier disparait du gabarit et l'application reste utilisable normalement.
    """
    from django.conf import settings
    contexte = {'google_client_id': getattr(settings, 'GOOGLE_OAUTH_CLIENT_ID', '')}
    contexte.update(extra)
    return contexte


def loginview(request):
    if request.user.is_authenticated:
        messages.success(request, "Bienvenue à AUTO-PIECE")
        request.session['show_ecom_order_guide'] = True
        return _redirect_after_login(request.user, request)
    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        password = request.POST.get("password", "")
        if not username or not password:
            messages.error(request, "Veuillez renseigner vos identifiants.")
            return render(request, "page/login.html", _auth_context())

        pending = CustomUser.objects.filter(username=username).first()
        if pending and pending.check_password(password) and not pending.is_active:
            messages.warning(
                request,
                "Votre compte n'est pas encore activé. Consultez l'email d'activation ou demandez un nouveau lien.",
            )
            return render(request, "page/login.html", _auth_context(show_resend=True))

        user = authenticate(request, username=username, password=password)
        if user is not None:
            if not user.is_active:
                messages.warning(request, "Votre compte n'est pas encore activé.")
                return render(request, "page/login.html", _auth_context(show_resend=True))
            login(request, user)
            civilite = "Mme" if user.genre == "Femme" else "Mr"
            messages.success(request, f"Bienvenue {civilite} {user.username}")
            return _redirect_after_login(user, request)
        messages.error(request, "Identifiant ou mot de passe incorrect.")
    return render(request, "page/login.html", _auth_context())

def Deconnexion(request):
    user = request.user
    role = getattr(user, 'role', None)
    is_superuser = getattr(user, 'is_superuser', False)
    username = getattr(user, 'username', '') or 'utilisateur'
    logout(request)
    messages.success(request, f'Vous êtes deconnecté {username}')

    # Clients → boutique ecom ; personnel magasin / admin → site holding
    if role == 'client' and not is_superuser:
        return redirect('ecom_index')
    return redirect('pbholdingsite')

OTP_EXPIRY_SECONDS = 300
def _get_active_otp_request(user):
    return PWD_FORGET.objects.filter(
        user_id=user, status='0', purpose=PWD_FORGET.PURPOSE_PASSWORD_RESET,
    ).order_by('-creat_at').first()

def _otp_remaining_seconds(reset_request):
    if not reset_request:
        return 0
    elapsed = (timezone.now() - reset_request.creat_at).total_seconds()
    return max(0, int(OTP_EXPIRY_SECONDS - elapsed))

def _otp_page_context(request):
    email = request.session.get('reset_email', '')
    if not email:
        return {
            'email': '',
            'remaining_seconds': 0,
            'otp_expired': True,
        }
    try:
        user = CustomUser.objects.get(email__iexact=email)
    except CustomUser.DoesNotExist:
        return {
            'email': email,
            'remaining_seconds': 0,
            'otp_expired': True,
        }
    reset_request = _get_active_otp_request(user)
    remaining = _otp_remaining_seconds(reset_request)
    return {
        'email': email,
        'remaining_seconds': remaining,
        'otp_expired': reset_request is None or remaining <= 0,
    }


class ForgotPasswordView(View):
    def get(self, request):
        return render(request, 'page/forgot_password.html')


class RequestEmailView(View):
    def post(self, request):
        form = PasswordResetRequestForm(request.POST)
        if not form.is_valid():
            messages.error(request, "Veuillez saisir une adresse email valide.")
            return render(request, 'page/forgot_password.html', {'form': form})

        email = form.cleaned_data['email'].strip().lower()
        try:
            user = CustomUser.objects.get(email__iexact=email)
            PWD_FORGET.objects.filter(
                user_id=user, status='0', purpose=PWD_FORGET.PURPOSE_PASSWORD_RESET,
            ).update(status='1')
            otp = random.randint(100000, 999999)
            PWD_FORGET.objects.create(
                user_id=user,
                otp=otp,
                status='0',
                purpose=PWD_FORGET.PURPOSE_PASSWORD_RESET,
            )
            request.session['reset_email'] = user.email
            if send_password_reset_otp_email(user, otp):
                messages.success(
                    request,
                    f"Un code à 6 chiffres a été envoyé à {user.email} (valide 5 min).",
                )
            else:
                messages.warning(request, "Le code a été généré mais l'email n'a pas pu être envoyé.")
            return redirect('otp')
        except CustomUser.DoesNotExist:
            messages.success(
                request,
                "Si un compte existe avec cet email, un code de vérification a été envoyé.",
            )
            return redirect('otp')


class OptValid(View):
    def get(self, request):
        if not request.session.get('reset_email'):
            messages.error(request, "Session expirée. Recommencez la procédure.")
            return redirect('forgot')
        context = _otp_page_context(request)
        if context['otp_expired']:
            messages.error(request, "Code OTP expiré. Demandez un nouveau code.")
        return render(request, 'page/otp.html', context)

    def post(self, request):
        otp_raw = request.POST.get('otp', '').strip()
        email = request.session.get('reset_email')
        if not otp_raw or not email:
            messages.error(request, "Session expirée. Recommencez la procédure.")
            return redirect('forgot')

        context = _otp_page_context(request)
        if context['otp_expired']:
            messages.error(request, "Code OTP expiré. Demandez un nouveau code.")
            return render(request, 'page/otp.html', context)

        try:
            otp = int(otp_raw)
            user = CustomUser.objects.get(email__iexact=email)
            reset_request = PWD_FORGET.objects.filter(
                user_id=user, otp=otp, status='0',
                purpose=PWD_FORGET.PURPOSE_PASSWORD_RESET,
            ).order_by('-creat_at').first()
            if not reset_request:
                messages.error(request, "Code OTP invalide.")
                return render(request, 'page/otp.html', context)
            if _otp_remaining_seconds(reset_request) <= 0:
                messages.error(request, "Code OTP expiré. Demandez un nouveau code.")
                context['otp_expired'] = True
                context['remaining_seconds'] = 0
                return render(request, 'page/otp.html', context)
            request.session['otp'] = otp
            request.session['otp_verified'] = True
            return redirect('verify_otp')
        except (ValueError, CustomUser.DoesNotExist):
            messages.error(request, "Code OTP invalide.")
            return render(request, 'page/otp.html', context)


class VerifyOtpView(View):
    def get(self, request):
        if not request.session.get('otp_verified'):
            messages.error(request, "Veuillez d'abord vérifier votre code OTP.")
            return redirect('otp')
        return render(request, 'page/reinitialise.html')

    def post(self, request):
        if not request.session.get('otp_verified'):
            messages.error(request, "Session expirée.")
            return redirect('forgot')

        form = PasswordResetConfirmForm(request.POST)
        if not form.is_valid():
            for errors in form.errors.values():
                for error in errors:
                    messages.error(request, error)
            return render(request, 'page/reinitialise.html', {'form': form})

        otp = request.session.get('otp')
        email = request.session.get('reset_email')
        try:
            user = CustomUser.objects.get(email__iexact=email)
            reset_request = PWD_FORGET.objects.get(
                user_id=user, otp=otp, status='0',
                purpose=PWD_FORGET.PURPOSE_PASSWORD_RESET,
            )
            if _otp_remaining_seconds(reset_request) <= 0:
                messages.error(request, "Code OTP expiré.")
                return redirect('forgot')
            user.set_password(form.cleaned_data['new_password'])
            user.save(update_fields=['password'])
            reset_request.status = '1'
            reset_request.save(update_fields=['status'])
            for key in ('otp', 'otp_verified', 'reset_email'):
                request.session.pop(key, None)
            messages.success(request, 'Mot de passe réinitialisé avec succès. Connectez-vous.')
            return redirect('connexion')
        except (CustomUser.DoesNotExist, PWD_FORGET.DoesNotExist):
            messages.error(request, "Session invalide. Recommencez la procédure.")
            return redirect('forgot')

class PasswordChangeView(LoginRequiredMixin, PasswordChangeView):
    login_url = 'connexion'
    form_class = ChangePasswordForm
    template_name = 'change_password.html'
    success_message = "Mot de passe réinitialisé avec succès👍✓✓"
    error_message = "Erreur de saisie ✘✘"
    success_url = reverse_lazy('password_change')
    def form_valid(self, form):
        reponse = super().form_valid(form)
        messages.success(self.request, self.success_message)
        return reponse
    def form_invalid(self, form):
        reponse = super().form_invalid(form)
        messages.error(self.request, self.error_message)
        return reponse 
    def get(self, request, *args, **kwargs):
        form = self.get_form()
        return render(request, self.template_name, {'form': form})
    
class PasswordChangeDoneView(View):
    def get(self, request):
         return render(request, 'password_change_done.html')

@login_required(login_url='connexion')
def change_password_profile(request):
    """Profil connecté : édition des infos + coordonnées + mot de passe + permissions."""
    user = request.user
    profil, _ = ProfilUser.objects.get_or_create(user=user)

    custom_permissions = user.custom_permissions.select_related('categorie').all()
    django_permissions = user.user_permissions.all()
    group_permissions = set()
    for group in user.groups.all():
        group_permissions.update(group.permissions.all())
    group_permissions = list(group_permissions)

    active_tab = request.POST.get('active_tab') or request.GET.get('tab') or 'info'
    if active_tab not in {'info', 'coords', 'security', 'permissions'}:
        active_tab = 'info'

    user_form = ProfileSelfUserForm(instance=user)
    profil_form = ProfileSelfProfilForm(instance=profil)
    password_form = ChangePasswordForm(user=user)

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'save_profile':
            user_form = ProfileSelfUserForm(request.POST, instance=user)
            profil_form = ProfileSelfProfilForm(request.POST, request.FILES, instance=profil)
            if user_form.is_valid() and profil_form.is_valid():
                user_form.save()
                profil_form.save()
                messages.success(request, 'Vos informations ont été enregistrées.')
                return redirect(f"{reverse('change_password')}?tab={active_tab}")
            messages.error(request, 'Veuillez corriger les erreurs du formulaire.')
            if active_tab not in {'info', 'coords'}:
                active_tab = 'info'
        elif action == 'change_password':
            active_tab = 'security'
            password_form = ChangePasswordForm(user=user, data=request.POST)
            if password_form.is_valid():
                password_form.save()
                update_session_auth_hash(request, password_form.user)
                messages.success(request, 'Votre mot de passe a été modifié avec succès.')
                return redirect(f"{reverse('change_password')}?tab=security")
            messages.error(request, 'Veuillez corriger les erreurs du formulaire.')

    display_name = user.get_full_name().strip() or user.username
    has_custom_photo = _user_has_custom_photo(profil)
    photo_url = ''
    if has_custom_photo:
        try:
            photo_url = profil.photo.url
        except Exception:
            has_custom_photo = False

    context = {
        'user': user,
        'profil': profil,
        'user_form': user_form,
        'profil_form': profil_form,
        'password_form': password_form,
        'custom_permissions': custom_permissions,
        'django_permissions': django_permissions,
        'group_permissions': group_permissions,
        'permissions_count': (
            custom_permissions.count()
            + django_permissions.count()
            + len(group_permissions)
        ),
        'active_tab': active_tab,
        'display_name': display_name,
        'user_initials': _user_nav_initials(user),
        'has_custom_photo': has_custom_photo,
        'photo_url': photo_url,
        'compte_actif': bool(user.is_active and getattr(profil, 'statut', 'actif') == 'actif'),
    }
    return render(request, 'page/change_pwd.html', context)

def ActivateCompte(request, pk):
    try:
        comptes = get_object_or_404(CustomUser, id=pk)
        comptes.is_active = True
        comptes.save()
        messages.success(request, f"Le compte de {comptes.username} a été activé.")
    except Exception as e:
        messages.error(request, f"Erreur lors de l'activation : {str(e)}")
    return redirect(request.META.get('HTTP_REFERER', 'default_view_name'))  
# fallback vers une vue si REFERER est vide

def DeactivatCompte(request, pk):
    try:
        comptes = get_object_or_404(CustomUser, id=pk)
        comptes.is_active = False
        comptes.save()
        messages.success(request, f"Le compte de {comptes.username} a été désactivé.")
    except Exception as e:
        messages.error(request, f"Erreur lors de la désactivation : {str(e)}")
    return redirect(request.META.get('HTTP_REFERER', 'default_view_name')) 

def delete_compte(request, pk):
    try:
        user = get_object_or_404(CustomUser, id=pk)
        user.delete()
        messages.success(request, f"Compte de {user.username} à été supprimé avec succès.")
    except Exception as e:
        messages.error(request, f"Erreur lors de la désactivation : {str(e)}")
    return redirect(request.META.get('HTTP_REFERER', 'default_view_name'))

@login_required(login_url='connexion')
@superuser_required
def update_permissions(request, pk):
    user = get_object_or_404(
        CustomUser.objects.select_related('local_entrepot').prefetch_related('custom_permissions'),
        id=pk,
    )
    
    if request.method == 'POST':
        permission_form = UserPermissionForm(request.POST)
        if permission_form.is_valid():
            permissions = permission_form.cleaned_data.get('permissions', [])
            if permissions:
                user.custom_permissions.set(permissions)
            else:
                user.custom_permissions.clear()
            messages.success(request, f'Permissions de {user.username} mises à jour avec succès.')
            
            if request.headers.get("x-requested-with") == "XMLHttpRequest":
                return JsonResponse({
                    "success": True,
                    "message": f'Permissions de {user.username} mises à jour avec succès.'
                })
            return redirect('add_compte')
        else:
            if request.headers.get("x-requested-with") == "XMLHttpRequest":
                html = render_to_string("partials/compte_perms.html", {
                    "permission_form": permission_form,
                    "object": user
                }, request=request)
                return JsonResponse({"success": False, "html": html})
    else:
        # Pré-remplir le formulaire avec les permissions actuelles de l'utilisateur
        initial_permissions = user.custom_permissions.all()
        permission_form = UserPermissionForm(initial={'permissions': initial_permissions})
    
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        html = render_to_string("partials/compte_perms.html", {
            "permission_form": permission_form,
            "object": user
        }, request=request)
        return HttpResponse(html)
    
    return render(request, "partials/compte_perms.html", {
        "permission_form": permission_form,
        "object": user
    })

# ============================================================
# VUES POUR LA GESTION DES PERMISSIONS PERSONNALISÉES
# ============================================================
@login_required(login_url='connexion')
@superuser_required
def list_permissions(request):
    """Liste toutes les permissions avec recherche, filtrage et pagination."""
    from django.core.paginator import Paginator

    permissions = CustomPermission.objects.all().select_related('categorie').order_by(
        'categorie__categorie', 'name',
    )

    search_query = request.GET.get('search', '')
    if search_query:
        permissions = permissions.filter(
            Q(name__icontains=search_query)
            | Q(url__icontains=search_query)
            | Q(categorie__categorie__icontains=search_query)
        )

    category_filter = request.GET.get('category', '')
    if category_filter:
        permissions = permissions.filter(categorie_id=category_filter)

    categories = TypeCustomPermission.objects.all().order_by('categorie')

    paginator = Paginator(permissions, 20)
    page_obj = paginator.get_page(request.GET.get('page'))

    context = {
        'permissions': page_obj.object_list,
        'page_obj': page_obj,
        'paginator': paginator,
        'categories': categories,
        'search_query': search_query,
        'category_filter': category_filter,
        'permissions_total': paginator.count,
    }
    return render(request, 'page/list_permissions.html', context)

@login_required(login_url='connexion')
@superuser_required
def create_permission(request):
    """Créer une nouvelle permission"""
    if request.method == 'POST':
        form = CustomPermissionForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ Permission créée avec succès!')
            if request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest":
                return JsonResponse({
                    "success": True,
                    "message": "Permission créée avec succès!",
                    "redirect_url": reverse("list_permissions")
                })
            return redirect('list_permissions')
        else:
            if request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest":
                html = render_to_string("partials/permission_form.html", {
                    "form": form
                }, request=request)
                return JsonResponse({"success": False, "html": html})
    else:
        form = CustomPermissionForm()
    
    # Pour GET, retourner le formulaire (avec ou sans AJAX)
    is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if is_ajax:
        html = render_to_string("partials/permission_form.html", {
            "form": form
        }, request=request)
        return HttpResponse(html)
    
    return render(request, "partials/permission_form.html", {"form": form})

@login_required(login_url='connexion')
@superuser_required
def update_permission(request, pk):
    """Modifier une permission existante"""
    permission = get_object_or_404(CustomPermission, pk=pk)
    
    if request.method == 'POST':
        form = CustomPermissionForm(request.POST, instance=permission)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ Permission mise à jour avec succès!')
            is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
            if is_ajax:
                return JsonResponse({
                    "success": True,
                    "message": "Permission mise à jour avec succès!",
                    "redirect_url": reverse("list_permissions")
                })
            return redirect('list_permissions')
        else:
            is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
            if is_ajax:
                html = render_to_string("partials/permission_form.html", {
                    "form": form,
                    "object": permission
                }, request=request)
                return JsonResponse({"success": False, "html": html})
    else:
        form = CustomPermissionForm(instance=permission)
    # Pour GET, retourner le formulaire (avec ou sans AJAX)
    is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if is_ajax:
        html = render_to_string("partials/permission_form.html", {
            "form": form,
            "object": permission
        }, request=request)
        return HttpResponse(html)
    return render(request, "partials/permission_form.html", {
        "form": form,
        "object": permission
    })

@login_required(login_url='connexion')
@superuser_required
def delete_permission(request, pk):
    """Supprimer une permission"""
    permission = get_object_or_404(CustomPermission, pk=pk)
    
    if request.method == 'POST':
        permission_name = permission.name
        permission.delete()
        messages.success(request, f'✅ Permission "{permission_name}" supprimée avec succès!')
        
        if request.headers.get("x-requested-with") == "XMLHttpRequest":
            return JsonResponse({
                "success": True,
                "message": f'Permission "{permission_name}" supprimée avec succès!'
            })
        return redirect('list_permissions')
    
    # Pour GET, retourner la confirmation
    html = render_to_string("partials/delete_permission_confirm.html", {
        "object": permission
    }, request=request)
    return HttpResponse(html)

@login_required(login_url='connexion')
def list_categories(request):
    """Liste toutes les catégories de permissions"""
    from django.db.models import Avg

    categories = TypeCustomPermission.objects.annotate(
        num_permissions=models.Count('cat_permis')
    ).order_by('categorie')
    permissions_total = CustomPermission.objects.count()
    avg_row = categories.aggregate(avg=Avg('num_permissions'))
    permissions_avg = round(avg_row['avg'] or 0, 1)
    context = {
        'categories': categories,
        'permissions_total': permissions_total,
        'permissions_avg': permissions_avg,
    }
    # Vérifier si c'est une requête AJAX
    is_ajax = (
        request.headers.get("x-requested-with") == "XMLHttpRequest" or 
        request.headers.get("X-Requested-With") == "XMLHttpRequest" or
        request.GET.get('ajax') == '1'
    )
    # Si c'est AJAX, retourner juste le contenu (sans layout complet)
    if is_ajax:
        return render(request, 'page/list_categories.html', context)
    # Sinon, retourner la page complète avec layout
    return render(request, 'page/list_categories_standalone.html', context)

@login_required(login_url='connexion')
def create_category(request):
    """Créer une nouvelle catégorie"""
    if request.method == 'POST':
        form = TypeCustomPermissionForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ Catégorie créée avec succès!')
            is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
            if is_ajax:
                return JsonResponse({
                    "success": True,
                    "message": "Catégorie créée avec succès!",
                    "redirect_url": reverse("list_perm_categories")
                })
            return redirect('list_perm_categories')
        else:
            is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
            if is_ajax:
                html = render_to_string("partials/category_form.html", {
                    "form": form
                }, request=request)
                return JsonResponse({"success": False, "html": html})
    else:
        form = TypeCustomPermissionForm()
    # Pour GET, retourner le formulaire (avec ou sans AJAX)
    is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if is_ajax:
        html = render_to_string("partials/category_form.html", {
            "form": form
        }, request=request)
        return HttpResponse(html)
    
    return render(request, "partials/category_form.html", {"form": form})

@login_required(login_url='connexion')
def update_category(request, pk):
    """Modifier une catégorie existante"""
    category = get_object_or_404(TypeCustomPermission, pk=pk)
    
    if request.method == 'POST':
        form = TypeCustomPermissionForm(request.POST, instance=category)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ Catégorie mise à jour avec succès!')
            is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
            if is_ajax:
                return JsonResponse({
                    "success": True,
                    "message": "Catégorie mise à jour avec succès!",
                    "redirect_url": reverse("list_perm_categories")
                })
            return redirect('list_perm_categories')
        else:
            is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
            if is_ajax:
                html = render_to_string("partials/category_form.html", {
                    "form": form,
                    "object": category
                }, request=request)
                return JsonResponse({"success": False, "html": html})
    else:
        form = TypeCustomPermissionForm(instance=category)
    
    # Pour GET, retourner le formulaire (avec ou sans AJAX)
    is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if is_ajax:
        html = render_to_string("partials/category_form.html", {
            "form": form,
            "object": category
        }, request=request)
        return HttpResponse(html)
    
    return render(request, "partials/category_form.html", {
        "form": form,
        "object": category
    })

@login_required(login_url='connexion')
def delete_category(request, pk):
    """Supprimer une catégorie"""
    category = get_object_or_404(TypeCustomPermission, pk=pk)
    if request.method == 'POST':
        category_name = category.categorie
        category.delete()
        messages.success(request, f'✅ Catégorie "{category_name}" supprimée avec succès!')
        is_ajax = request.headers.get("x-requested-with") == "XMLHttpRequest" or request.headers.get("X-Requested-With") == "XMLHttpRequest"
        if is_ajax:
            return JsonResponse({
                "success": True,
                "message": f'Catégorie "{category_name}" supprimée avec succès!'
            })
        return redirect('list_perm_categories')
    # Pour GET, retourner la confirmation
    html = render_to_string("partials/delete_category_confirm.html", {
        "object": category
    }, request=request)
    return HttpResponse(html)

@login_required(login_url='connexion')
@superuser_required
def import_permissions_excel(request):
    """Importer les permissions depuis un fichier Excel"""
    if request.method == 'POST':
        form = ImportPermissionsForm(request.POST, request.FILES)
        if form.is_valid():
            excel_file = request.FILES['excel_file']
            update_existing = form.cleaned_data.get('update_existing', False)
            
            try:
                # Sauvegarder le fichier temporairement
                with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
                    for chunk in excel_file.chunks():
                        tmp_file.write(chunk)
                    tmp_file_path = tmp_file.name
                
                # Charger le fichier Excel
                workbook = load_workbook(tmp_file_path, data_only=True)
                sheet = workbook.active
                
                created_count = 0
                updated_count = 0
                skipped_count = 0
                error_count = 0
                
                with transaction.atomic():
                    for row_num in range(2, sheet.max_row + 1):  # Ignorer la première ligne (en-têtes)
                        try:
                            name = sheet.cell(row=row_num, column=1).value
                            categorie_name = sheet.cell(row=row_num, column=2).value
                            url = sheet.cell(row=row_num, column=3).value
                            
                            if not name or not url:
                                skipped_count += 1
                                continue
                            
                            name = str(name).strip()
                            url = str(url).strip()
                            categorie_name = str(categorie_name).strip() if categorie_name else 'Autres'
                            
                            # Créer ou récupérer la catégorie
                            category, _ = TypeCustomPermission.objects.get_or_create(
                                categorie=categorie_name,
                                defaults={'categorie': categorie_name}
                            )
                            
                            # Créer ou mettre à jour la permission
                            if update_existing:
                                perm, created = CustomPermission.objects.update_or_create(
                                    url=url,
                                    defaults={
                                        'name': name,
                                        'categorie': category
                                    }
                                )
                                if created:
                                    created_count += 1
                                else:
                                    updated_count += 1
                            else:
                                if CustomPermission.objects.filter(url=url).exists():
                                    skipped_count += 1
                                else:
                                    CustomPermission.objects.create(
                                        name=name,
                                        categorie=category,
                                        url=url
                                    )
                                    created_count += 1
                        
                        except Exception as e:
                            error_count += 1
                            continue
                
                # Supprimer le fichier temporaire
                os.unlink(tmp_file_path)
                
                messages.success(
                    request,
                    f'✅ Importation terminée! Créées: {created_count}, '
                    f'{"Mises à jour: " + str(updated_count) + ", " if update_existing else ""}'
                    f'Ignorées: {skipped_count}, Erreurs: {error_count}'
                )
                return redirect('list_permissions')
            
            except Exception as e:
                messages.error(request, f'❌ Erreur lors de l\'importation: {str(e)}')
    else:
        form = ImportPermissionsForm()
    
    context = {
        'form': form,
    }
    return render(request, 'page/import_permissions.html', context)

@login_required(login_url='connexion')
@superuser_required
def export_permissions_excel(request):
    """Exporter les permissions vers un fichier Excel"""
    permissions = CustomPermission.objects.all().select_related('categorie').order_by('categorie__categorie', 'name')
    
    if not permissions.exists():
        messages.warning(request, '⚠️ Aucune permission à exporter.')
        return redirect('list_permissions')
    
    # Créer le classeur Excel
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Permissions'
    
    # Style pour l'en-tête
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    
    # En-têtes
    headers = ['Nom de la Permission', 'Catégorie', 'URL/Name']
    for col_num, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Largeur des colonnes
    sheet.column_dimensions['A'].width = 40
    sheet.column_dimensions['B'].width = 25
    sheet.column_dimensions['C'].width = 50
    
    # Données
    for row_num, permission in enumerate(permissions, start=2):
        sheet.cell(row=row_num, column=1).value = permission.name
        sheet.cell(row=row_num, column=2).value = permission.categorie.categorie
        sheet.cell(row=row_num, column=3).value = permission.url
    
    # Créer la réponse HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'permissions_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    workbook.save(response)
    return response


def register_view(request):
    if request.user.is_authenticated:
        messages.info(request, "Vous êtes déjà connecté.")
        return redirect('connexion')
    if request.method == 'POST':
        form = ClientRegistrationForm(request.POST)
        if form.is_valid():
            try:
                with transaction.atomic():
                    user = form.save()
                    ProfilUser.objects.get_or_create(user=user)
                    email_sent = _send_account_activation(user, request)
                if email_sent:
                    messages.success(
                        request,
                        f"Compte créé ! Un email d'activation a été envoyé à {user.email}. "
                        "Cliquez sur le lien pour activer votre compte, puis connectez-vous.",
                    )
                else:
                    messages.warning(
                        request,
                        "Compte créé, mais l'email d'activation n'a pas pu être envoyé. "
                        "Utilisez « Renvoyer l'activation ».",
                    )
                return redirect('connexion')
            except Exception as e:
                messages.error(request, f"Erreur lors de la création du compte : {e}")
        else:
            for errors in form.errors.values():
                for error in errors:
                    messages.error(request, error)
    else:
        form = ClientRegistrationForm()
    return render(request, 'page/register.html', _auth_context(
        form=form,
        google_label="S'inscrire avec Google",
        google_context='signup',
    ))


def activate_account_view(request, token):
    try:
        verification_token = EmailVerificationToken.objects.select_related('user').get(token=token)
        user = verification_token.user
        if not verification_token.is_valid():
            if verification_token.used and user.is_active:
                messages.info(request, "Votre compte est déjà activé. Vous pouvez vous connecter.")
            else:
                messages.error(
                    request,
                    "Ce lien d'activation a expiré (validité 24 h). "
                    "Demandez un nouveau lien depuis la page « Renvoyer l'activation ».",
                )
                return redirect('resend-activation')
            return redirect('connexion')
        with transaction.atomic():
            user.is_active = True
            user.save(update_fields=['is_active'])
            verification_token.used = True
            verification_token.used_at = timezone.now()
            verification_token.save(update_fields=['used', 'used_at'])
        messages.success(
            request,
            f"Félicitations {user.get_full_name() or user.username} ! Votre compte est activé. Connectez-vous.",
        )
        return redirect('connexion')
    except EmailVerificationToken.DoesNotExist:
        messages.error(request, "Lien d'activation invalide.")
        return redirect('connexion')


def verify_email_view(request, token):
    return activate_account_view(request, token)


def resend_activation_view(request):
    if request.user.is_authenticated:
        messages.info(request, "Vous êtes déjà connecté.")
        return redirect('connexion')

    if request.method == 'POST':
        form = ResendActivationForm(request.POST)
        if form.is_valid():
            email = form.cleaned_data['email'].strip().lower()
            try:
                user = CustomUser.objects.get(email__iexact=email)
                if user.is_active:
                    messages.info(request, "Ce compte est déjà activé. Vous pouvez vous connecter.")
                    return redirect('connexion')
                _invalidate_activation_tokens(user)
                token_obj = EmailVerificationToken.objects.create(user=user)
                if send_verification_email(user, token_obj.token, request):
                    messages.success(request, f"Un nouvel email d'activation a été envoyé à {user.email}.")
                else:
                    messages.error(request, "L'email n'a pas pu être envoyé. Réessayez plus tard.")
                return redirect('connexion')
            except CustomUser.DoesNotExist:
                messages.error(request, "Aucun compte associé à cet email.")
    else:
        form = ResendActivationForm()
    return render(request, 'page/resend_activation.html', {'form': form})

# ---------------------------------------------------------------------------
# Localités / entrepôts (LocalEntrepot + CreneauDisponibilite)
# ---------------------------------------------------------------------------
def _localites_map_payload(localites):
    """Payload JSON pour la carte Leaflet (marqueurs + infos popup)."""
    payload = []
    for loc in localites:
        creneaux = [
            {
                'jour': c.get_jour_display(),
                'ouverture': c.heure_ouverture.strftime('%H:%M'),
                'fermeture': c.heure_fermeture.strftime('%H:%M'),
                'statut': c.statut,
            }
            for c in loc.creneaux.all()
        ]
        payload.append({
            'code': str(loc.code),
            'nom': loc.nom,
            'lat': float(loc.latitude) if loc.latitude is not None else None,
            'lng': float(loc.longitude) if loc.longitude is not None else None,
            'ouvert': bool(loc.statut),
            'statut': loc.statut_ouverture,
            'creneaux': creneaux,
            'edit_url': reverse('update_localite', args=[loc.code]),
            'delete_url': reverse('delete_localite', args=[loc.code]),
        })
    return payload


def _localite_form_context(form, creneau_formset, local=None):
    return {
        'form': form,
        'creneau_formset': creneau_formset,
        'localite': local,
        'is_edit': local is not None and local.pk is not None,
    }


@login_required(login_url='connexion')
def add_localite(request):
    """Liste des localités + création (formulaire modal avec carte + créneaux)."""
    localites = LocalEntrepot.objects.prefetch_related('creneaux').order_by('nom')
    open_modal = False

    if request.method == 'POST':
        form = LocalEntrepotForm(request.POST)
        creneau_formset = CreneauDisponibiliteFormSet(request.POST, instance=LocalEntrepot())
        if form.is_valid() and creneau_formset.is_valid():
            try:
                with transaction.atomic():
                    local = form.save()
                    creneau_formset.instance = local
                    creneau_formset.save()
                messages.success(request, f'Localité « {local.nom} » créée avec succès.')
                return redirect('add_localite')
            except Exception as e:
                messages.error(request, f'Erreur lors de la création : {e}')
                open_modal = True
        else:
            for error in form.errors.values():
                messages.error(request, error)
            for err in creneau_formset.non_form_errors():
                messages.error(request, err)
            for f in creneau_formset:
                for field_errors in f.errors.values():
                    for err in field_errors:
                        messages.error(request, err)
            open_modal = True
    else:
        form = LocalEntrepotForm()
        creneau_formset = CreneauDisponibiliteFormSet(instance=LocalEntrepot())

    context = {
        'form': form,
        'creneau_formset': creneau_formset,
        'localites': localites,
        'localites_json': json.dumps(_localites_map_payload(localites), ensure_ascii=False),
        'open_modal': open_modal,
        'is_edit': False,
    }
    return render(request, 'page/add_localite.html', context)


@login_required(login_url='connexion')
def update_localite(request, pk):
    """Modifier une localité + ses créneaux (AJAX form ou POST classique)."""
    local = get_object_or_404(
        LocalEntrepot.objects.prefetch_related('creneaux'),
        pk=pk,
    )
    is_ajax = (
        request.headers.get('x-requested-with') == 'XMLHttpRequest'
        or request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    )

    if request.method == 'POST':
        form = LocalEntrepotForm(request.POST, instance=local)
        creneau_formset = CreneauDisponibiliteFormSet(request.POST, instance=local)
        if form.is_valid() and creneau_formset.is_valid():
            with transaction.atomic():
                form.save()
                creneau_formset.save()
            messages.success(request, f'Localité « {local.nom} » mise à jour.')
            if is_ajax:
                return JsonResponse({
                    'success': True,
                    'message': f'Localité « {local.nom} » mise à jour.',
                    'redirect_url': reverse('add_localite'),
                })
            return redirect('add_localite')
        if is_ajax:
            html = render_to_string(
                'partials/localite_form.html',
                _localite_form_context(form, creneau_formset, local),
                request=request,
            )
            return JsonResponse({'success': False, 'html': html})
        for error in form.errors.values():
            messages.error(request, error)
        for f in creneau_formset:
            for field_errors in f.errors.values():
                for err in field_errors:
                    messages.error(request, err)
        return redirect('add_localite')

    form = LocalEntrepotForm(instance=local)
    creneau_formset = CreneauDisponibiliteFormSet(instance=local)
    html = render_to_string(
        'partials/localite_form.html',
        _localite_form_context(form, creneau_formset, local),
        request=request,
    )
    if is_ajax:
        return JsonResponse({
            'success': True,
            'html': html,
            'nom': local.nom,
            'lat': float(local.latitude) if local.latitude is not None else None,
            'lng': float(local.longitude) if local.longitude is not None else None,
        })
    return HttpResponse(html)


@login_required(login_url='connexion')
@require_POST
def delete_localite(request, pk):
    """Supprimer une localité depuis la page liste."""
    local = get_object_or_404(LocalEntrepot, pk=pk)
    nom = local.nom
    try:
        local.delete()
        messages.success(request, f'Localité « {nom} » supprimée.')
    except Exception as e:
        messages.error(request, f'Impossible de supprimer « {nom} » : {e}')
    return redirect('add_localite')
