"""Helpers partagés pour exports Excel / PDF des listes magasin."""
from io import BytesIO

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from xhtml2pdf import pisa

from .period_filters import append_query_string


HEADER_FILL = PatternFill(start_color='EA580C', end_color='EA580C', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)


def export_urls(request, excel_name, pdf_name, url_kwargs=None):
    """URLs Excel/PDF avec query string du filtre courant."""
    kw = url_kwargs or {}
    excel = reverse(excel_name, kwargs=kw)
    pdf = reverse(pdf_name, kwargs=kw)
    return {
        'page_export_excel_url': append_query_string(excel, request),
        'page_export_pdf_url': append_query_string(pdf, request),
        'page_export_excel_title': 'Exporter en Excel',
        'page_export_pdf_title': 'Exporter en PDF',
    }


def excel_response(filename, headers, rows, sheet_title='Export'):
    """Génère une réponse HTTP Excel (openpyxl)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            if value is None:
                value = ''
            ws.cell(row=r_idx, column=c_idx, value=value)

    for col in ws.columns:
        letter = col[0].column_letter
        max_len = 10
        for cell in col:
            try:
                max_len = max(max_len, min(len(str(cell.value or '')), 40))
            except Exception:
                pass
        ws.column_dimensions[letter].width = max_len + 2

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def pdf_response(request, *, title, headers, rows, filename, subtitle=''):
    """Génère une réponse HTTP PDF via template générique + pisa."""
    html = render_to_string(
        'mag/exports/export_table_pdf.html',
        {
            'title': title,
            'subtitle': subtitle,
            'headers': headers,
            'rows': rows,
        },
        request=request,
    )
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    status = pisa.CreatePDF(html, dest=response)
    if status.err:
        return HttpResponse('Erreur lors de la génération du PDF.', status=500)
    return response


def cell(value, default='—'):
    if value is None or value == '':
        return default
    return value
